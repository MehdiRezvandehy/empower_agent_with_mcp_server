from pathlib import Path
import json

import fitz
from docx import Document
from openpyxl import load_workbook

from mcp.server.fastmcp import FastMCP


mcp = FastMCP("Document Search MCP")


# =========================================================
# CONFIGURATION
# =========================================================

IGNORED_DIRECTORIES = {
    ".git",
    ".venv",
    "venv",
    "__pycache__",
    "node_modules",
    ".idea",
    ".vscode",
}


TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".log",
    ".py", ".pyw", ".js", ".ts",
    ".java", ".c", ".cpp", ".h", ".hpp",
    ".r", ".sql",
    ".json", ".xml", ".html", ".htm",
    ".csv", ".tsv",
    ".ipynb",
}


SUPPORTED_EXTENSIONS = (
    TEXT_EXTENSIONS
    | {
        ".pdf",
        ".docx",
        ".xlsx",
        ".xlsm",
    }
)


# =========================================================
# HELPERS
# =========================================================

def should_ignore(path: Path):

    return any(
        part in IGNORED_DIRECTORIES
        for part in path.parts
    )


# =========================================================
# READ TEXT
# =========================================================

def read_text_file(path: Path):

    try:

        with open(
            path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as f:

            return f.read()

    except Exception:

        return ""


# =========================================================
# READ PDF
# =========================================================

def read_pdf_file(path: Path):

    text_parts = []

    try:

        with fitz.open(path) as document:

            for page_number, page in enumerate(
                document,
                start=1
            ):

                text = page.get_text().strip()

                if text:

                    text_parts.append(
                        f"[PAGE {page_number}]\n{text}"
                    )

    except Exception:

        return ""

    return "\n\n".join(text_parts)


# =========================================================
# READ DOCX
# =========================================================

def read_docx_file(path: Path):

    text_parts = []

    try:

        document = Document(path)

        # Paragraphs

        for paragraph in document.paragraphs:

            text = paragraph.text.strip()

            if text:

                text_parts.append(text)

        # Tables

        for table in document.tables:

            for row in table.rows:

                values = []

                for cell in row.cells:

                    value = cell.text.strip()

                    if value:

                        values.append(value)

                if values:

                    text_parts.append(
                        " | ".join(values)
                    )

    except Exception:

        return ""

    return "\n".join(text_parts)


# =========================================================
# READ EXCEL
# =========================================================

def read_excel_file(path: Path):

    text_parts = []

    try:

        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True
        )

        for worksheet in workbook.worksheets:

            text_parts.append(
                f"[SHEET: {worksheet.title}]"
            )

            for row in worksheet.iter_rows(
                values_only=True
            ):

                values = [
                    str(value)
                    for value in row
                    if value is not None
                ]

                if values:

                    text_parts.append(
                        " | ".join(values)
                    )

        workbook.close()

    except Exception:

        return ""

    return "\n".join(text_parts)


# =========================================================
# READ JUPYTER NOTEBOOK
# =========================================================

def read_ipynb_file(path: Path):

    try:

        with open(
            path,
            "r",
            encoding="utf-8",
            errors="ignore"
        ) as f:

            notebook = json.load(f)

    except Exception:

        return ""

    text_parts = []

    for cell_number, cell in enumerate(
        notebook.get("cells", []),
        start=1
    ):

        source = "".join(
            cell.get("source", [])
        )

        text_parts.append(
            f"[CELL {cell_number}]\n{source}"
        )

    return "\n".join(text_parts)


# =========================================================
# GENERIC FILE READER
# =========================================================

def read_file(path: Path):

    extension = path.suffix.lower()

    if extension in TEXT_EXTENSIONS:

        if extension == ".ipynb":

            return read_ipynb_file(path)

        return read_text_file(path)

    if extension == ".pdf":

        return read_pdf_file(path)

    if extension == ".docx":

        return read_docx_file(path)

    if extension in {".xlsx", ".xlsm"}:

        return read_excel_file(path)

    return ""


# =========================================================
# MCP SEARCH TOOL
# =========================================================

@mcp.tool()
def search_documents(
    directory: str,
    query: str,
    max_results: int = 20,
) -> str:

    """
    Recursively search filenames and file contents.

    Searches supported documents including:

    PDF
    DOCX
    XLSX
    CSV
    TXT
    Markdown
    Python
    JSON
    XML
    HTML
    Jupyter notebooks
    and other text-based files.

    Returns JSON containing:

    - folders_searched
    - files_searched
    - number_of_results
    - matching documents
    - surrounding text context
    """

    root = Path(directory)

    # =====================================================
    # VALIDATE DIRECTORY
    # =====================================================

    if not root.exists():

        return json.dumps({
            "error": f"Directory does not exist: {directory}"
        })

    if not root.is_dir():

        return json.dumps({
            "error": f"Not a directory: {directory}"
        })


    # =====================================================
    # VALIDATE QUERY
    # =====================================================

    query = query.strip()

    if not query:

        return json.dumps({
            "error": "Query cannot be empty."
        })

    query_lower = query.lower()


    # =====================================================
    # COUNTERS
    # =====================================================

    folders_searched = 0

    files_searched = 0

    results = []


    # =====================================================
    # RECURSIVE SEARCH
    # =====================================================

    for path in root.rglob("*"):

        # -------------------------------------------------
        # IGNORE UNWANTED DIRECTORIES
        # -------------------------------------------------

        if should_ignore(path):

            continue


        # -------------------------------------------------
        # COUNT FOLDERS
        # -------------------------------------------------

        if path.is_dir():

            folders_searched += 1

            continue


        # -------------------------------------------------
        # ONLY PROCESS FILES
        # -------------------------------------------------

        if not path.is_file():

            continue


        # -------------------------------------------------
        # SUPPORTED FILE TYPES
        # -------------------------------------------------

        extension = path.suffix.lower()

        if extension not in SUPPORTED_EXTENSIONS:

            continue


        # -------------------------------------------------
        # COUNT FILE
        # -------------------------------------------------

        files_searched += 1


        # =================================================
        # FILENAME SEARCH
        # =================================================

        filename_match = (
            query_lower in path.name.lower()
        )

        if filename_match:

            results.append({

                "file": str(path),

                "file_name": path.name,

                "file_type": extension,

                "match_type": "filename",

                "context":
                    f"Filename contains: {query}"

            })

            if len(results) >= max_results:

                break


        # =================================================
        # CONTENT SEARCH
        # =================================================

        text = read_file(path)

        if not text:

            continue


        position = text.lower().find(
            query_lower
        )

        if position == -1:

            continue


        # =================================================
        # CONTEXT
        # =================================================

        context_size = 1000

        start = max(
            0,
            position - context_size
        )

        end = min(
            len(text),
            position + len(query) + context_size
        )

        context = text[start:end]


        # =================================================
        # ADD RESULT
        # =================================================

        results.append({

            "file": str(path),

            "file_name": path.name,

            "file_type": extension,

            "match_type": "content",

            "context": context

        })


        if len(results) >= max_results:

            break


    # =========================================================
    # RETURN STRUCTURED JSON
    # =========================================================

    response = {

        "query": query,

        "directory": directory,

        "folders_searched": folders_searched,

        "files_searched": files_searched,

        "number_of_results": len(results),

        "results": results

    }


    return json.dumps(
        response,
        indent=2,
        ensure_ascii=False
    )


# =========================================================
# START SERVER
# =========================================================

if __name__ == "__main__":
    mcp.settings.host = "0.0.0.0"
    mcp.settings.port = 8000

    mcp.run(
        transport="streamable-http"
    )