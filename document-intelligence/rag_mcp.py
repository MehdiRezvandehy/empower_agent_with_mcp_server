from pathlib import Path
import hashlib

import chromadb
import fitz
from docx import Document
from openpyxl import load_workbook
from openai import OpenAI

from mcp.server.fastmcp import FastMCP


mcp = FastMCP("Document RAG MCP")


# ============================================================
# CONFIGURATION
# ============================================================

TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".log",
    ".py", ".pyw", ".js", ".ts",
    ".java", ".c", ".cpp", ".h", ".hpp",
    ".r", ".sql",
    ".json", ".xml", ".html", ".htm",
    ".csv", ".tsv",
    ".ipynb",
}


CHROMA_PATH = "./chroma_db"

COLLECTION_NAME = "documents"


# ============================================================
# READ DOCUMENTS
# ============================================================

def read_documents_for_rag(directory):

    root = Path(directory)

    if not root.exists():
        raise ValueError(
            f"Directory does not exist: {directory}"
        )

    documents = []

    for path in root.rglob("*"):

        if not path.is_file():
            continue

        extension = path.suffix.lower()

        # ----------------------------------------------------
        # TEXT FILES
        # ----------------------------------------------------

        if extension in TEXT_EXTENSIONS:

            try:

                with open(
                    path,
                    "r",
                    encoding="utf-8",
                    errors="ignore"
                ) as f:

                    text = f.read()

            except Exception:

                continue

        # ----------------------------------------------------
        # PDF
        # ----------------------------------------------------

        elif extension == ".pdf":

            try:

                text_parts = []

                with fitz.open(path) as pdf:

                    for page_number, page in enumerate(
                        pdf,
                        start=1
                    ):

                        page_text = page.get_text()

                        if page_text.strip():

                            text_parts.append(
                                f"[PAGE {page_number}]\n"
                                f"{page_text}"
                            )

                text = "\n\n".join(
                    text_parts
                )

            except Exception:

                continue

        # ----------------------------------------------------
        # DOCX
        # ----------------------------------------------------

        elif extension == ".docx":

            try:

                document = Document(path)

                text_parts = []

                for paragraph in document.paragraphs:

                    if paragraph.text.strip():

                        text_parts.append(
                            paragraph.text
                        )

                text = "\n".join(
                    text_parts
                )

            except Exception:

                continue

        # ----------------------------------------------------
        # EXCEL
        # ----------------------------------------------------

        elif extension in {".xlsx", ".xlsm"}:

            try:

                workbook = load_workbook(
                    path,
                    read_only=True,
                    data_only=True
                )

                text_parts = []

                for worksheet in workbook.worksheets:

                    text_parts.append(
                        f"[SHEET: {worksheet.title}]"
                    )

                    for row in worksheet.iter_rows(
                        values_only=True
                    ):

                        values = [
                            str(value)
                            for value in row
                            if value is not None
                        ]

                        if values:

                            text_parts.append(
                                " | ".join(values)
                            )

                workbook.close()

                text = "\n".join(
                    text_parts
                )

            except Exception:

                continue

        else:

            continue

        # ----------------------------------------------------
        # SAVE DOCUMENT
        # ----------------------------------------------------

        if text and text.strip():

            documents.append(
                {
                    "text": text,
                    "metadata": {
                        "source": str(path),
                        "file_name": path.name,
                        "file_type": extension
                    }
                }
            )

    return documents


# ============================================================
# CREATE CHUNKS
# ============================================================

def create_chunks(
    documents,
    chunk_size=1000,
    chunk_overlap=200
):

    chunks = []

    for document in documents:

        text = document["text"]

        metadata = document["metadata"]

        text = text.strip()

        if not text:
            continue

        start = 0

        while start < len(text):

            end = start + chunk_size

            chunk_text = text[
                start:end
            ].strip()

            if chunk_text:

                chunks.append(
                    {
                        "text": chunk_text,
                        "metadata": metadata.copy()
                    }
                )

            start += (
                chunk_size -
                chunk_overlap
            )

    return chunks


# ============================================================
# CREATE CHUNK ID
# ============================================================

def create_chunk_id(
    chunk,
    index
):

    source = chunk["metadata"].get(
        "source",
        ""
    )

    text = chunk["text"]

    value = (
        f"{source}|"
        f"{index}|"
        f"{text}"
    )

    return hashlib.md5(
        value.encode("utf-8")
    ).hexdigest()


# ============================================================
# ADD CHUNKS TO CHROMA
# ============================================================

def add_chunks(
    chunks,
    api_key,
    embedding_model
):

    if not chunks:

        return 0

    openai_client = OpenAI(
        api_key=api_key
    )

    chroma_client = chromadb.PersistentClient(
        path=CHROMA_PATH
    )

    collection = chroma_client.get_or_create_collection(
        name=COLLECTION_NAME
    )

    texts = []

    metadatas = []

    ids = []

    for i, chunk in enumerate(chunks):

        texts.append(
            chunk["text"]
        )

        metadatas.append(
            chunk["metadata"]
        )

        ids.append(
            create_chunk_id(
                chunk,
                i
            )
        )

    response = openai_client.embeddings.create(
        model=embedding_model,
        input=texts
    )

    embeddings = [
        item.embedding
        for item in response.data
    ]

    collection.upsert(
        ids=ids,
        documents=texts,
        metadatas=metadatas,
        embeddings=embeddings
    )

    return len(chunks)


# ============================================================
# CREATE QUERY EMBEDDING
# ============================================================

def create_embedding(
    text,
    api_key,
    embedding_model
):

    client = OpenAI(
        api_key=api_key
    )

    response = client.embeddings.create(
        model=embedding_model,
        input=[text]
    )

    return response.data[0].embedding


# ============================================================
# RETRIEVE DOCUMENTS
# ============================================================

def retrieve_documents(
    question,
    api_key,
    embedding_model,
    n_results=5
):

    embedding = create_embedding(
        question,
        api_key,
        embedding_model
    )

    chroma_client = chromadb.PersistentClient(
        path=CHROMA_PATH
    )

    collection = chroma_client.get_or_create_collection(
        name=COLLECTION_NAME
    )

    results = collection.query(
        query_embeddings=[embedding],
        n_results=n_results
    )

    return results


# ============================================================
# GENERATE RAG ANSWER
# ============================================================

def generate_rag_answer(
    question,
    documents,
    metadatas,
    api_key,
    model_name="gpt-4.1-mini",
    temperature=0
):

    client = OpenAI(
        api_key=api_key
    )

    # --------------------------------------------------------
    # BUILD CONTEXT
    # --------------------------------------------------------

    context_parts = []

    for i, (
        document,
        metadata
    ) in enumerate(
        zip(
            documents,
            metadatas
        ),
        start=1
    ):

        source = metadata.get(
            "source",
            "Unknown source"
        )

        context_parts.append(
            f"""
SOURCE {i}:
{source}

CONTENT:
{document}
"""
        )

    context = "\n\n".join(
        context_parts
    )

    # --------------------------------------------------------
    # PROMPT
    # --------------------------------------------------------

    prompt = f"""
You are a document question-answering assistant.

Answer the user's question using ONLY the information
provided in the document context below.

If the answer cannot be found in the documents,
say that the information was not found.

Always identify the source document when possible.

USER QUESTION:
{question}

DOCUMENT CONTEXT:
{context}
"""

    # --------------------------------------------------------
    # OPENAI LLM
    # --------------------------------------------------------

    response = client.chat.completions.create(
        model=model_name,
        temperature=temperature,
        messages=[
            {
                "role": "system",
                "content": (
                    "You answer questions using retrieved "
                    "document context."
                )
            },
            {
                "role": "user",
                "content": prompt
            }
        ]
    )

    return response.choices[0].message.content


# ============================================================
# MCP RAG TOOL
# ============================================================

@mcp.tool()
def run_rag(
    directory: str,
    question: str,
    api_key: str,
    embedding_model: str,
    model_name: str,
    chunk_size: int = 1000,
    chunk_overlap: int = 200,
    temperature: float = 0,
    n_results: int = 5
) -> dict:

    """
    Read documents from a directory, create embeddings,
    store them in Chroma, retrieve relevant chunks, and
    generate an answer using an OpenAI model.

    Parameters:

    directory:
        Directory containing documents.

    question:
        User's question.

    api_key:
        OpenAI API key.

    embedding_model:
        OpenAI embedding model.

    model_name:
        OpenAI chat model.

    chunk_size:
        Number of characters per chunk.

    chunk_overlap:
        Number of overlapping characters.

    temperature:
        LLM temperature.

    n_results:
        Number of chunks to retrieve.
    """

    # ========================================================
    # VALIDATION
    # ========================================================

    if not directory:

        return {
            "error": "Directory cannot be empty."
        }

    if not question:

        return {
            "error": "Question cannot be empty."
        }

    if not api_key:

        return {
            "error": "OpenAI API key cannot be empty."
        }

    # ========================================================
    # STEP 1 — READ DOCUMENTS
    # ========================================================

    documents = read_documents_for_rag(
        directory
    )

    if not documents:

        return {
            "error": (
                "No supported documents "
                "were found."
            )
        }

    # ========================================================
    # STEP 2 — CREATE CHUNKS
    # ========================================================

    chunks = create_chunks(
        documents,
        chunk_size=chunk_size,
        chunk_overlap=chunk_overlap
    )

    # ========================================================
    # STEP 3 — CREATE EMBEDDINGS
    # ========================================================

    chunks_added = add_chunks(
        chunks,
        api_key=api_key,
        embedding_model=embedding_model
    )

    # ========================================================
    # STEP 4 — RETRIEVE
    # ========================================================

    retrieved = retrieve_documents(
        question=question,
        api_key=api_key,
        embedding_model=embedding_model,
        n_results=n_results
    )

    retrieved_documents = (
        retrieved.get("documents", [[]])[0]
    )

    retrieved_metadata = (
        retrieved.get("metadatas", [[]])[0]
    )

    if not retrieved_documents:

        return {
            "error": (
                "No relevant information "
                "was found."
            )
        }

    # ========================================================
    # STEP 5 — GENERATE ANSWER
    # ========================================================

    answer = generate_rag_answer(
        question=question,
        documents=retrieved_documents,
        metadatas=retrieved_metadata,
        api_key=api_key,
        model_name=model_name,
        temperature=temperature
    )

    # ========================================================
    # RETURN RESULT
    # ========================================================

    sources = []

    for metadata in retrieved_metadata:

        sources.append(
            metadata.get(
                "source",
                "Unknown"
            )
        )

    return {
        "answer": answer,
        "sources": sources,
        "documents_found": len(documents),
        "chunks_created": len(chunks),
        "chunks_added": chunks_added,
        "chunks_retrieved": len(
            retrieved_documents
        )
    }


# ============================================================
# SERVER
# ============================================================

if __name__ == "__main__":
    mcp.settings.host = "0.0.0.0"
    mcp.settings.port = 8001
    mcp.settings.transport_security.allowed_hosts = [
        "rag-server:8001",
        "localhost:8001",
        "127.0.0.1:8001"
    ]

    mcp.run(
        transport="streamable-http"
    )